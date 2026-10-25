"""協力業者（解体屋さん・造成屋さん）への見積依頼書。

金額は入れず、数量・仕様・現場条件・写真・図面をまとめて「この数量で見積もってください」と渡す資料。
PDF（依頼書一式）と Excel（単価を記入すると金額・合計が出る数量表）を作る。
"""
import datetime
import os

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Spacer

from . import docpages as D
from . import elevation, grading_draw, quantities
from .common import MM

KIND_LABEL = {"demolition": "解体工事", "grading": "造成工事"}
STRUCT = {"wood": "木造", "light_steel": "軽量鉄骨造", "heavy_steel": "重量鉄骨造", "rc": "RC造"}


def _lines(kind, demo, gest):
    """見積を依頼する項目（金額なし）。区分ごと。"""
    est = demo if kind == "demolition" else gest
    out = []
    for ln in est["lines"]:
        name = ln["name"]
        if ln["cat"] == "手続" and "滅失登記" in name:
            continue  # 滅失登記は土地家屋調査士の業務なので依頼しない
        basis = ln.get("basis", "")
        if ln["cat"] == "建物本体" and "解体" in name and "外階段" not in name:
            basis = "上屋・基礎の撤去、分別解体、廃材の運搬・処分、足場・養生、簡易整地まで。" + basis.split("推定")[0] + "推定"
        out.append({"cat": ln["cat"], "name": name, "qty": ln["qty"], "unit": ln["unit"], "basis": basis})
    if kind == "demolition":
        for ln in demo.get("asbestos_extra") or []:
            out.append({"cat": "アスベスト（含有時）", "name": ln["name"].replace(" ※含有判明時", ""), "qty": ln["qty"],
                        "unit": ln["unit"], "basis": "事前調査で含有が判明した場合の単価をご記入ください"})
    return out


def _overview(spec, kind, demo, grad):
    b = spec.get("building") or {}
    acc = spec.get("access") or {}
    fin = b.get("finishes") or {}
    rows = []
    if kind == "demolition" and demo:
        q = demo["quantities"]
        rows += [
            ("構造・階数", f"{STRUCT.get(b.get('structure'), b.get('structure', ''))} {b.get('floors', '')}階建" + (f"（{b.get('use')}）" if b.get("use") else "")),
            ("延床面積（推定）", f"約{q['floor_area_m2']:.1f}㎡（{q['floor_area_tsubo']:.1f}坪）／建築面積 約{q['building_area_m2']:.1f}㎡"),
            ("外形寸法（推定）", f"約{b.get('footprint', {}).get('width_m', 0):.2f}m × {b.get('footprint', {}).get('depth_m', 0):.2f}m、軒高 約{b.get('eave_h', 0):.2f}m"),
            ("推定築年", f"{b.get('built_year_est')}年頃（アスベスト事前調査は未実施）"),
        ]
        rows += [(f"{k}（推定）", f"{v}　{note}" if note else v) for k, v, note in quantities.summary_rows(q)[1:]]
        if q.get("units_count") and q["units_count"] > 1:
            rows.append(("戸数・間取り", f"{q['units_count']}戸（{q.get('units_layout', '')}）"))
        if fin:
            rows.append(("外部仕上げ（推定）", "、".join(f"{k}：{v}" for k, v in fin.items() if v)))
        annex = [a["name"] for a in q.get("annexes", [])]
        if annex:
            rows.append(("付属物", "、".join(annex)))
        lo = demo["leftover"]
        rows.append(("残置物（推定）", f"約{lo['volume_m3']}㎥（家財が残る戸数 {lo['occupied_units']}戸、屋外 {lo['extra_m3']}㎥）"))
    if grad:
        rows.append(("敷地面積（推定）", f"約{grad['area_m2']:.0f}㎡（{grad['area_m2'] / 3.30579:.1f}坪）"))
        rows.append(("敷地の高低差（推定）", f"敷地内 約{grad['existing_diff']:.2f}m"))
    road = (spec.get("site") or {}).get("road") or {}
    rows.append(("前面道路", f"幅員 約{acc.get('road_width_m') or road.get('width_m', '')}m" + (f"（{road.get('name')}）" if road.get("name") else "")))
    if acc:
        if acc.get("machine"):
            rows.append(("重機・車両（推定）", f"重機：{acc.get('machine')}／車両：{acc.get('truck', '')}"))
        if acc.get("neighbor_gap_m"):
            rows.append(("隣地との離れ（推定）", f"最も近いところで 約{acc['neighbor_gap_m']}m"))
        if acc.get("carry_distance_m"):
            rows.append(("搬出距離（推定）", f"車両から 約{acc['carry_distance_m']}m"))
        if acc.get("notes"):
            rows.append(("搬出・近隣の条件", acc["notes"]))
    return rows


def _requests(kind):
    if kind == "demolition":
        return [
            "数量は現地写真からの推定です。現地確認のうえ、数量に過不足があれば修正してお見積りください。",
            "アスベスト事前調査（書面・目視・分析）の費用と、含有していた場合の除去単価を分けてご記入ください。",
            "地中埋設物（古い基礎・浄化槽・ガラ等）が出た場合の単価または扱いをご記入ください。",
            "産業廃棄物の運搬・処分費、重機回送費、養生・足場、近隣挨拶、道路使用許可の扱い（含む／別途）をご明記ください。",
            "建設リサイクル法の届出代行の可否、工期の目安（着工〜完了）をご記入ください。",
            "上記の数量表にない作業が必要な場合は、空欄の行に追記してください。",
        ]
    return [
        "数量（土量・擁壁延長・高さ）は写真と公開データ（国土地理院の標高）からの推定です。現地確認・測量のうえ修正してお見積りください。",
        "残土の受入先と処分単価、購入土の単価・土質をご記入ください。",
        "擁壁は現場打ち／プレキャストのどちらで見積もったかと、構造計算・確認申請の扱い（含む／別途）をご明記ください。",
        "盛土規制法・開発許可の申請が必要な場合の費用と期間の目安をご記入ください。",
        "地盤改良が必要になった場合の概算、工期の目安をご記入ください。",
        "上記の数量表にない作業が必要な場合は、空欄の行に追記してください。",
    ]


def build_rfq_pdf(path, spec, base_dir, demo, grad, gest, kind, rfq, annotate):
    def render(target, total):
        c = rl_canvas.Canvas(target, pagesize=A4)
        c.setTitle(f"{KIND_LABEL[kind]} 見積依頼書")
        doc = D.Doc(c, spec, total, header=f"{KIND_LABEL[kind]} 見積依頼書")
        S = D.st()
        # 1. 依頼書（1ページ目）＋ 物件概要
        p = spec.get("project", {})
        doc.a4()
        y = D.PH - 32
        doc.text(D.ML, y, "　　　　　　　　　　　　　　　　御中", 12)
        c.setStrokeColor(D.INK)
        c.setLineWidth(0.6)
        c.line(D.ML * MM, (y - 2) * MM, (D.ML + 82) * MM, (y - 2) * MM)
        doc.text(D.PW - D.MR, y + 2, f"依頼日　{p.get('date', '')}", 9, D.INK, "r")
        yy = y - 5
        for line in [rfq.get("company") or D.COMPANY, rfq.get("contact", "")]:
            if line:
                doc.text(D.PW - D.MR, yy, line, 9, D.INK, "r")
                yy -= 5
        y -= 24
        doc.text(D.PW / 2, y, f"{KIND_LABEL[kind]}　見積依頼書", 20, D.INK, "c")
        c.setStrokeColor(D.GOLD)
        c.setLineWidth(1.2)
        c.line((D.PW / 2 - 12) * MM, (y - 5) * MM, (D.PW / 2 + 12) * MM, (y - 5) * MM)
        y -= 14
        intro = D.P(f"平素より大変お世話になっております。下記の{KIND_LABEL[kind]}について、お見積りをお願いいたします。"
                    "次ページ以降の数量表の単価・金額欄にご記入のうえ、ご返送ください（同封の Excel に単価を入力すると、金額と合計が自動で計算されます）。",
                    S["body"])
        flows = [intro, Spacer(1, 3 * MM)]
        due = rfq.get("due") or "　　　　年　　月　　日"
        head = [("件名", f"{p.get('name', '')}　{KIND_LABEL[kind]}"), ("工事場所", p.get("address", "")),
                ("回答期限", due), ("工期の希望", rfq.get("period") or "ご提案ください"),
                ("現地確認", "可能です（事前にご連絡ください）"),
                ("添付資料", "数量表、現地写真、配置図・立面図" + ("・造成計画図" if kind == "grading" else "") + "（いずれも推定を含みます）")]
        flows.append(D.table([[D.P(k, S["cell_sub"]), D.P(v, S["cell"])] for k, v in head], [28, 146], header=False))
        flows += [Spacer(1, 5 * MM), D.P("物件概要（写真・公開データからの推定を含みます）", S["h2"]),
                  D.table([[D.P(k, S["cell_sub"]), D.P(v, S["cell"])] for k, v in _overview(spec, kind, demo, grad)], [38, 136], header=False)]
        doc.flow(flows, None, None, y=y)

        # 2. 数量表（単価・金額は空欄）
        rows = [[D.hcell("No."), D.hcell("項目・仕様"), D.hcell("数量"), D.hcell("単位"), D.hcell("単価"), D.hcell("金額")]]
        styles = []
        n = 0
        cats = []
        items = _lines(kind, demo, gest)
        for it in items:
            if it["cat"] not in cats:
                cats.append(it["cat"])
        for cat in cats:
            rows.append(["", D.P(cat, S["cell"]), "", "", "", ""])
            styles += [("BACKGROUND", (0, len(rows) - 1), (-1, len(rows) - 1), D.SOFT), ("SPAN", (1, len(rows) - 1), (-1, len(rows) - 1))]
            for it in [x for x in items if x["cat"] == cat]:
                n += 1
                rows.append([D.P(str(n), S["cell_sub"]), D.name_cell(it["name"], it["basis"] or None, S), D.P(D.qty_s(it["qty"]), S["cell_r"]),
                             D.P(it["unit"], S["cell"]), "", ""])
        for _ in range(4):
            n += 1
            rows.append([D.P(str(n), S["cell_sub"]), D.P("（追加項目）", S["cell_sub"]), "", "", "", ""])
        for label in ("小計", "諸経費", "合計（税別）", "消費税", "合計（税込）"):
            rows.append(["", D.P(label, D.ParagraphStyle("r", parent=S["cell"], alignment=2)), "", "", "", ""])
            r = len(rows) - 1
            styles.append(("SPAN", (1, r), (4, r)))
            if label == "小計":
                styles.append(("LINEABOVE", (0, r), (-1, r), 0.8, D.INK))
            if label.startswith("合計"):
                styles.append(("BACKGROUND", (5, r), (5, r), D.GOLD_SOFT))
        styles += [("GRID", (4, 1), (5, -1), 0.35, D.LINE)]
        t = D.table(rows, [10, 92, 18, 12, 20, 22], styles=styles)
        notes = [Spacer(1, 4 * MM), D.P("お見積りの際のお願い", S["h2"])] + [D.P("・" + s, S["bullet"]) for s in _requests(kind)]
        doc.flow([t] + notes, "QUANTITIES", f"{KIND_LABEL[kind]} 数量表（ご記入用）")

        # 3. 写真台帳・図面
        if spec.get("photos"):
            D.photo_pages(doc, spec, base_dir, annotate)
        if grad:
            grading_draw.render_plan(c, spec, grad, title="現況平面図・解体範囲図（参考）" if kind == "demolition" else "現況平面図（参考）",
                                     sheet_no="B-01", mode="existing")
        if kind == "demolition" and demo:
            elevation.render_elevations(c, spec)
        if kind == "grading" and grad:
            grading_draw.render_plan(c, spec, grad)
            grading_draw.render_sections(c, spec, grad)
        c.save()
        return doc.page

    import io
    total = render(io.BytesIO(), None)
    render(path, total)


def write_rfq_xlsx(path, spec, demo, gest, grad, kind, rfq):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    FJ = "游ゴシック"
    ink, sub, line, soft, gold_soft = "333333", "7A7F85", "CDD6DD", "F3F5F6", "F6F1E4"
    thin = Side(style="thin", color=line)
    YEN = '#,##0;[Red]-#,##0'
    p = spec.get("project", {})
    wb = Workbook()
    ws = wb.active
    ws.title = "数量表（ご記入用）"
    for i, w in enumerate([6, 46, 10, 7, 13, 15, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "&P / &N"
    ws.merge_cells("A1:G1")
    ws["A1"] = f"{KIND_LABEL[kind]}　見積依頼 数量表"
    ws["A1"].font = Font(name=FJ, size=16, bold=True, color=ink)
    ws.row_dimensions[1].height = 30
    info = [("件名", p.get("name", "")), ("工事場所", p.get("address", "")), ("依頼元", " ".join(x for x in [rfq.get("company") or D.COMPANY, rfq.get("contact", "")] if x)),
            ("依頼日", p.get("date", "")), ("回答期限", rfq.get("due") or ""), ("御社名", "")]
    r = 2
    for k, v in info:
        ws.cell(r, 1, k).font = Font(name=FJ, size=9, color=sub)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        c = ws.cell(r, 2, v)
        c.font = Font(name=FJ, size=10, color=ink)
        if k == "御社名":
            c.fill = PatternFill("solid", fgColor="FFFDF2")
            c.border = Border(bottom=thin)
        r += 1
    r += 1
    ws.cell(r, 1, "黄色の欄（単価・諸経費）をご記入ください。金額と合計は自動で計算されます。数量は写真からの推定のため、現地確認のうえ修正してください。").font = Font(name=FJ, size=9, color=sub)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    r += 1
    for i, h in enumerate(["No.", "項目・仕様", "数量", "単位", "単価", "金額", "備考"], 1):
        c = ws.cell(r, i, h)
        c.font = Font(name=FJ, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ink)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = ws.cell(r + 1, 1)
    r += 1
    fill_in = PatternFill("solid", fgColor="FFFDF2")
    first = r
    n = 0
    items = _lines(kind, demo, gest)
    cats = []
    for it in items:
        if it["cat"] not in cats:
            cats.append(it["cat"])
    for cat in cats + ["追加項目"]:
        ws.cell(r, 2, cat).font = Font(name=FJ, size=10, bold=True, color=ink)
        for col in range(1, 8):
            ws.cell(r, col).fill = PatternFill("solid", fgColor=soft)
        r += 1
        rows = [x for x in items if x["cat"] == cat] if cat != "追加項目" else [{"name": "", "qty": None, "unit": "", "basis": ""}] * 4
        for it in rows:
            n += 1
            ws.cell(r, 1, n).font = Font(name=FJ, size=9, color=sub)
            c = ws.cell(r, 2, it["name"])
            c.font = Font(name=FJ, size=10, color=ink)
            c.alignment = Alignment(wrap_text=True, vertical="center")
            if it["qty"] is not None:
                ws.cell(r, 3, it["qty"]).number_format = "#,##0.0#"
            ws.cell(r, 4, it["unit"])
            c = ws.cell(r, 5)
            c.fill, c.number_format = fill_in, YEN
            ws.cell(r, 6, f'=IF(E{r}="","",ROUND(C{r}*E{r},0))').number_format = YEN
            c = ws.cell(r, 7, it.get("basis", ""))
            c.font = Font(name=FJ, size=8.5, color=sub)
            c.alignment = Alignment(wrap_text=True, vertical="center")
            for col in range(1, 8):
                ws.cell(r, col).border = Border(bottom=thin)
            if cat == "追加項目":
                for col in (2, 3, 4):
                    ws.cell(r, col).fill = fill_in
            ws.row_dimensions[r].height = max(18, 14 * max(1, -(-len(it["name"]) // 22), -(-len(it.get("basis", "")) // 20)))
            r += 1
    last = r - 1
    r += 1
    ws.cell(r, 5, "小計").font = Font(name=FJ, size=10, bold=True)
    ws.cell(r, 6, f"=SUM(F{first}:F{last})").number_format = YEN
    ws.cell(r + 1, 5, "諸経費").font = Font(name=FJ, size=10)
    c = ws.cell(r + 1, 6)
    c.fill, c.number_format = fill_in, YEN
    ws.cell(r + 2, 5, "合計（税別）").font = Font(name=FJ, size=10, bold=True)
    ws.cell(r + 2, 6, f"=F{r}+N(F{r + 1})").number_format = YEN
    ws.cell(r + 3, 5, "消費税（10%）").font = Font(name=FJ, size=10)
    ws.cell(r + 3, 6, f"=ROUND(F{r + 2}*0.1,0)").number_format = YEN
    ws.cell(r + 4, 5, "合計（税込）").font = Font(name=FJ, size=11, bold=True)
    c = ws.cell(r + 4, 6, f"=F{r + 2}+F{r + 3}")
    c.number_format, c.font = YEN, Font(name=FJ, size=11, bold=True)
    for rr in (r, r + 2, r + 4):
        for col in (5, 6):
            ws.cell(rr, col).fill = PatternFill("solid", fgColor=gold_soft)
            ws.cell(rr, col).border = Border(top=Side(style="thin", color=ink))
    r += 6
    ws.cell(r, 1, "お見積りの際のお願い").font = Font(name=FJ, size=10, bold=True)
    r += 1
    for s in _requests(kind):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        c = ws.cell(r, 1, "・" + s)
        c.font = Font(name=FJ, size=9, color=ink)
        c.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[r].height = 26
        r += 1

    ov = wb.create_sheet("物件概要")
    ov.sheet_view.showGridLines = False
    ov.column_dimensions["A"].width = 22
    ov.column_dimensions["B"].width = 70
    ov["A1"] = "物件概要（写真・公開データからの推定を含みます）"
    ov["A1"].font = Font(name=FJ, size=13, bold=True)
    for i, (k, v) in enumerate(_overview(spec, kind, demo, grad), 3):
        ov.cell(i, 1, k).font = Font(name=FJ, size=9, color=sub)
        c = ov.cell(i, 2, v)
        c.font = Font(name=FJ, size=10)
        c.alignment = Alignment(wrap_text=True)
        ov.cell(i, 1).border = ov.cell(i, 2).border = Border(bottom=thin)
    wb.save(path)


def default_due(days=7):
    return (datetime.date.today() + datetime.timedelta(days=days)).strftime("%Y年%m月%d日")
