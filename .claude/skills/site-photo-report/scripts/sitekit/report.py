"""一式資料（PDF・Excel・DXF）の組み立て。"""
import datetime
import io
import os

from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Paragraph, Table, TableStyle

from . import demolition, docpages, elevation, grading, grading_draw
from .common import FONT, MM, ensure_font, font_path, man_yen, yen
from .drawing import MARGIN, SHEET_H, SHEET_W, frame, jp, new_sheet, text_at

CONF = {"確定": colors.HexColor("#1b7f3b"), "要確認": colors.HexColor("#b26a00"), "推測": colors.HexColor("#555555")}


def P(text, style):
    return Paragraph(jp(text), style)


def _styles():
    ensure_font()
    return {
        "h": ParagraphStyle("h", fontName=FONT, fontSize=15, leading=20),
        "b": ParagraphStyle("b", fontName=FONT, fontSize=10.5, leading=15),
        "L": ParagraphStyle("L", fontName=FONT, fontSize=12.5, leading=19),
        "s": ParagraphStyle("s", fontName=FONT, fontSize=8, leading=11),
        "t": ParagraphStyle("t", fontName=FONT, fontSize=8.5, leading=11),
    }


def _table(data, col_w, header=True, font_size=8.5, zebra=True, align_right_cols=()):
    st = _styles()
    ps = ParagraphStyle("c", parent=st["t"], fontSize=font_size, leading=font_size * 1.3)
    rows = [[Paragraph(jp(x), ps) if not isinstance(x, Paragraph) else x for x in r] for r in data]
    t = Table(rows, colWidths=[w * MM for w in col_w], repeatRows=1 if header else 0)
    style = [("FONT", (0, 0), (-1, -1), FONT, font_size), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 1.5),
             ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]
    if header:
        style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8ecf2")))
    if zebra:
        for i in range(1 if header else 0, len(rows)):
            if i % 2 == 0:
                style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#f7f7f7")))
    for col in align_right_cols:
        style.append(("ALIGN", (col, 1), (col, -1), "RIGHT"))
    t.setStyle(TableStyle(style))
    return t


def _draw_flow(c, flow, x_mm, y_top_mm, w_mm, h_mm):
    """flowable を上詰めで描く。収まらない分は分割して残りを返す。"""
    w, h = flow.wrapOn(c, w_mm * MM, h_mm * MM)
    if h <= h_mm * MM:
        flow.drawOn(c, x_mm * MM, y_top_mm * MM - h)
        return None, h / MM
    parts = flow.split(w_mm * MM, h_mm * MM)
    if not parts:
        flow.drawOn(c, x_mm * MM, y_top_mm * MM - h)
        return None, h / MM
    first = parts[0]
    w, h = first.wrapOn(c, w_mm * MM, h_mm * MM)
    first.drawOn(c, x_mm * MM, y_top_mm * MM - h)
    return (parts[1] if len(parts) > 1 else None), h / MM


def _page_with_flows(c, spec, title, sheet_no, flows, notes=None):
    """flowable を A3 図枠のページに流し込む（複数ページ可）。"""
    queue = list(flows)
    n = 0
    while queue:
        new_sheet(c)
        ax, ay, aw, ah = frame(c, title, f"{sheet_no}-{n + 1}" if n else sheet_no, spec.get("project", {}), notes=notes)
        ytop = ay + ah
        col_w = aw
        while queue and ytop - ay > 15:
            f = queue[0]
            if f == "COLBREAK":
                queue.pop(0)
                break
            rest, used = _draw_flow(c, f, ax, ytop, col_w, ytop - ay)
            queue.pop(0)
            if rest is not None:
                queue.insert(0, rest)
                break
            ytop -= used + 4
        c.showPage()
        n += 1
    return n


# ---------- 写真 ----------

def annotate_photo(path, marks, max_px=1400):
    """marks: [{"bbox":[x0,y0,x1,y1] (0-1), "label":"..."}] を描いた画像を返す。"""
    im = Image.open(path).convert("RGB")
    im.thumbnail((max_px, max_px))
    d = ImageDraw.Draw(im)
    fs = max(14, im.width // 45)
    font = ImageFont.truetype(font_path(), fs)
    for i, m in enumerate(marks or []):
        x0, y0, x1, y1 = m["bbox"]
        box = (x0 * im.width, y0 * im.height, x1 * im.width, y1 * im.height)
        col = (220, 30, 30) if m.get("color", "red") == "red" else (30, 90, 210)
        d.rectangle(box, outline=col, width=max(2, im.width // 300))
        label = f"{i + 1}. {m['label']}"
        tw = d.textlength(label, font=font)
        ty = box[1] - fs - 6 if box[1] > fs + 6 else box[3] + 2
        d.rectangle((box[0], ty, box[0] + tw + 8, ty + fs + 4), fill=col)
        d.text((box[0] + 4, ty + 1), label, fill=(255, 255, 255), font=font)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=85)
    buf.seek(0)
    return ImageReader(buf), im.size


def photo_pages(c, spec, base_dir):
    photos = spec.get("photos", [])
    if not photos:
        return
    per = 4
    for p0 in range(0, len(photos), per):
        new_sheet(c)
        ax, ay, aw, ah = frame(c, "写真台帳（読み取り箇所）", f"P-{p0 // per + 1:02d}", spec.get("project", {}))
        cells = [(ax, ay + ah / 2, aw / 2 - 3, ah / 2 - 2), (ax + aw / 2 + 3, ay + ah / 2, aw / 2 - 3, ah / 2 - 2),
                 (ax, ay, aw / 2 - 3, ah / 2 - 4), (ax + aw / 2 + 3, ay, aw / 2 - 3, ah / 2 - 4)]
        for ph, (x, y, w, h) in zip(photos[p0:p0 + per], cells):
            path = ph["path"] if os.path.isabs(ph["path"]) else os.path.join(base_dir, ph["path"])
            if not os.path.exists(path):
                text_at(c, x + 2, y + h / 2, f"写真が見つかりません: {ph['path']}", 2.5)
                continue
            img, (iw, ih) = annotate_photo(path, ph.get("marks"))
            cap_h = 4 + 3.3 * min(4, len(ph.get("marks", [])))
            bw, bh = w * 0.6, h - 8
            k = min(bw / iw, bh / ih)
            dw, dh = iw * k, ih * k
            c.drawImage(img, x * MM, (y + h - dh - 6) * MM, dw * MM, dh * MM)
            text_at(c, x, y + h - 4, f"写真{ph.get('no', '')} {ph.get('caption', '')}", 2.8)
            ty = y + h - 10
            for i, m in enumerate(ph.get("marks", [])):
                for line in _wrap(f"{i + 1}. {m['label']}" + (f" — {m['basis']}" if m.get("basis") else ""), 34):
                    text_at(c, x + dw + 3, ty, line, 2.2)
                    ty -= 3.3
        c.showPage()


def _wrap(s, n):
    out = []
    while len(s) > n:
        out.append(s[:n])
        s = s[n:]
    out.append(s)
    return out


# ---------- 各ページ ----------

def cover_page(c, spec, demo, grad, gest):
    st = _styles()
    new_sheet(c)
    p = spec.get("project", {})
    ax, ay, aw, ah = frame(c, "現地写真 調査レポート（表紙・総括）", "00", p)
    text_at(c, ax + 4, ay + ah - 12, p.get("name", "現地写真 調査レポート"), 8)
    text_at(c, ax + 4, ay + ah - 22, f"所在地：{p.get('address', '')}", 4)
    text_at(c, ax + 4, ay + ah - 29, f"作成日：{p.get('date', '')}　　依頼内容：{p.get('purpose', '解体概算・造成計画')}", 3.4)
    # 航空写真
    aer = spec.get("terrain", {}).get("aerial_photo")
    if aer and os.path.exists(aer):
        c.drawImage(aer, (ax + aw - 120) * MM, (ay + ah - 128) * MM, 118 * MM, 118 * MM, preserveAspectRatio=True)
        text_at(c, ax + aw - 120, ay + ah - 132, "航空写真（国土地理院 シームレス空中写真）", 2.3)
    rows = [["項目", "概算（標準）", "幅（安め〜高め）", "メモ"]]
    if demo:
        q = demo["quantities"]
        rows.append(["解体工事", man_yen(demo["total"]["mid"]), f"{man_yen(demo['total']['low'])} 〜 {man_yen(demo['total']['high'])}",
                     f"延床 約{q['floor_area_m2']:.0f}㎡（{q['floor_area_tsubo']:.1f}坪）／坪単価 約{demo['per_tsubo_mid'] / 10000:.1f}万円（諸経費込み）"])
        if demo["asbestos_extra"]:
            a = demo["asbestos_extra"][0]
            rows.append(["（アスベスト含有時の追加）", man_yen(a["mid"]), f"{man_yen(a['low'])} 〜 {man_yen(a['high'])}", "事前調査で含有が判明した場合のみ"])
    if gest:
        rows.append(["造成工事", man_yen(gest["total"]["mid"]), f"{man_yen(gest['total']['low'])} 〜 {man_yen(gest['total']['high'])}",
                     f"切土{grad['cut_m3']:.0f}㎥／盛土{grad['fill_m3']:.0f}㎥／擁壁 {sum(w['length_m'] for w in grad['walls'] if w['type'] != 'none'):.1f}m"])
    for a in (grad or {}).get("alternatives", [])[1:]:
        rows.append([f"（造成 別案）{a['label']}", man_yen(a["total_mid"]), "",
                     f"FH {a['fh_m'] - (grad['road']['center_avg'] if grad.get('road') else 0):+.2f}m／擁壁{a['wall_m']:.0f}m" + ("／道路より低い（排水・浸水注意）" if a["below_road"] else "")])
    if demo and gest:
        rows.append(["合計", man_yen(demo["total"]["mid"] + gest["total"]["mid"]),
                     f"{man_yen(demo['total']['low'] + gest['total']['low'])} 〜 {man_yen(demo['total']['high'] + gest['total']['high'])}", "税別"])
    t = _table(rows, [58, 36, 60, 100], align_right_cols=(1,), font_size=10.5)
    _, used = _draw_flow(c, t, ax + 4, ay + ah - 38, 262, 90)
    y = ay + ah - 38 - used - 8
    flows = [P("<b>この敷地の要点</b>", st["L"])]
    for s in spec.get("summary", []):
        flows.append(P("・" + s, st["L"]))
    flows.append(P("<b>精度について</b>", st["b"]))
    flows.append(P("本資料は現地写真と公開データ（国土地理院の標高・航空写真等）からの推定で作成した参考資料です。"
                           "寸法・数量は±10〜20%、標高は±0.2〜0.5m程度の誤差を含みます。見積は公開相場の単価幅（安め／標準／高め）による概算で、"
                           "正式見積・設計には現地調査・測量・地盤調査・アスベスト事前調査が必要です。", st["s"]))
    yy = y
    for f in flows:
        _, used = _draw_flow(c, f, ax + 4, yy, 265, yy - ay)
        yy -= used + 1.5
    c.showPage()


def site_info_page(c, spec, grad):
    st = _styles()
    info = spec.get("land_info", [])
    rows = [["項目", "内容", "確実性", "出典・根拠"]]
    for it in info:
        conf = it.get("confidence", "推測")
        rows.append([it["item"], it["value"], f"<font color='{CONF.get(conf, colors.black).hexval()}'>{conf}</font>"
                     if conf in CONF else conf, it.get("source", "")])
    if grad:
        r = grad.get("road") or {}
        rows.append(["敷地内の高低差", f"{grad['existing_diff']:.2f}m（最低 T.P.{grad['existing_min']:.2f} ／ 最高 T.P.{grad['existing_max']:.2f}）",
                     "推測", "国土地理院 1mDEM＋写真補正"])
        if r:
            rows.append(["前面道路", f"道路中心 T.P.{r['center_avg']:.2f}m、縦断勾配 約{abs(r['slope_pct']):.1f}%", "推測", "国土地理院 1mDEM"])
    flows = [_table(rows, [50, 190, 22, 120])]
    if spec.get("market_notes"):
        flows.append(P("<b>相場・土地勘メモ</b>", st["b"]))
        for s in spec["market_notes"]:
            flows.append(P("・" + s, st["b"]))
    _page_with_flows(c, spec, "土地の概況（公開データ）", "01", flows,
                     notes=["確実性：確定＝公的資料・条文で確認／要確認＝窓口・現地・原本で確かめるまで未確定／推測＝地図・写真・一般傾向からの見立て"])


def building_page(c, spec, demo):
    st = _styles()
    b = spec["building"]
    q = demo["quantities"]
    rows = [["項目", "推定値", "根拠（写真・基準物）"]]
    for r in b.get("measure_basis", []):
        rows.append([r["item"], r["value"], r.get("basis", "")])
    flows = [P("<b>建物の推定寸法と根拠</b>", st["b"]), _table(rows, [55, 60, 265])]
    rows2 = [["数量項目", "数量", "算出"],
             ["建築面積", f"{q['building_area_m2']:.1f} ㎡", "外形寸法から"],
             ["延床面積", f"{q['floor_area_m2']:.1f} ㎡（{q['floor_area_tsubo']:.1f}坪）", f"{b.get('floors')}階建"],
             ["外周長", f"{q['perimeter_m']:.1f} m", ""],
             ["足場架面積", f"{q['scaffold_m2']:.1f} ㎡", "(外周+8m)×(軒高+1m)"]]
    if q.get("units_count"):
        chk = q.get("units_area_check_m2")
        rows2.append(["戸数・間取り", f"{q['units_count']}戸（{q.get('units_layout', '')}）",
                      f"専有合計 約{chk:.0f}㎡（延床との差＝共用部・壁芯差）" if chk else ""])
    for a in q.get("annexes", []):
        rows2.append([a["name"], f"{a['area_m2']:.1f} ㎡", a.get("kind", "")])
    if q.get("asbestos_suspect_parts"):
        rows2.append(["アスベスト含有の可能性がある部位", "、".join(q["asbestos_suspect_parts"]), f"推定築年 {b.get('built_year_est')}年頃"])
    flows += [P("<b>解体数量（建物）</b>", st["b"]), _table(rows2, [55, 80, 245])]
    lo = demo["leftover"]
    if lo["furniture"]:
        rows3 = [["品目", "想定数量", "体積(㎥)", "家電リサイクル"]]
        for f in lo["furniture"]:
            rows3.append([f["name"], f"{f['qty']:g}", f"{f['m3']:.2f}", f.get("recycle") or ""])
        flows += [P(f"<b>残置物（家財）の推定</b>　{lo['layout']} × {lo['occupied_units']}戸、"
                            f"積込み体積 約{lo['volume_m3']}㎥（家財リスト換算 {lo['volume_from_list_m3']}㎥ ／ 相場表 {lo['volume_from_table_m3']}㎥ ／ 屋外 {lo['extra_m3']}㎥）", st["b"]),
                  _table(rows3, [80, 30, 30, 40], align_right_cols=(1, 2))]
    _page_with_flows(c, spec, "建物・数量の推定", "02", flows)


def estimate_page(c, spec, est, title, sheet_no, extra_lines=None, conditions=None):
    st = _styles()
    rows = [["区分", "項目", "数量", "単位", "単価（標準）", "金額（標準）", "安め", "高め", "算出根拠", "出典"]]
    for ln in est["lines"] + (extra_lines or []):
        p = ln.get("price") or {}
        rows.append([ln["cat"], ln["name"], f"{ln['qty']:,.1f}" if isinstance(ln["qty"], float) else f"{ln['qty']:,}",
                     ln["unit"], yen(p.get("mid")), yen(ln["mid"]), yen(ln["low"]), yen(ln["high"]),
                     ln.get("basis", ""), ",".join(ln.get("src", []))])
    flows = [_table(rows, [24, 82, 18, 12, 24, 28, 24, 24, 110, 22], align_right_cols=(2, 4, 5, 6, 7), font_size=7.8)]
    sub, over, tot = est["subtotal"], est["overhead"], est["total"]
    sm = [["", "安め", "標準", "高め"],
          ["小計（地域係数・条件割増 反映後）", yen(sub["low"]), yen(sub["mid"]), yen(sub["high"])],
          [f"諸経費（{est['overhead_rate']['low'] * 100:.0f}〜{est['overhead_rate']['high'] * 100:.0f}%）", yen(over["low"]), yen(over["mid"]), yen(over["high"])],
          ["合計（税別）", yen(tot["low"]), yen(tot["mid"]), yen(tot["high"])]]
    flows.append(_table(sm, [80, 40, 40, 40], align_right_cols=(1, 2, 3)))
    memo = [f"地域係数 ×{est.get('region_factor', 1.0)}"]
    if conditions:
        memo += conditions
    flows.append(P("適用：" + "／".join(memo) + "。単価の出典は巻末「出典一覧」。単価を自社単価に差し替える場合は同梱の Excel を使用。", st["s"]))
    if extra_lines:
        flows.append(P("※「アスベスト（含有時のみ）」の行は合計に含めていません。", st["s"]))
    _page_with_flows(c, spec, title, sheet_no, flows)


def notes_page(c, spec, demo, gest):
    st = _styles()
    flows = [P("<b>次に確認すること（見積・設計の精度を上げるため）</b>", st["b"])]
    for s in spec.get("next_checks", []):
        flows.append(P("□ " + s, st["b"]))
    flows.append(P("<b>出典一覧（単価・データ）</b>", st["b"]))
    rows = [["ID", "出典", "URL", "確認日"]]
    seen = set()
    for est in (demo, gest):
        if not est:
            continue
        for s in est.get("sources", []):
            if s["id"] in seen:
                continue
            seen.add(s["id"])
            rows.append([s["id"], s.get("title", ""), s.get("url", ""), s.get("accessed", "")])
    for s in spec.get("data_attribution", []):
        rows.append(["-", s, "", ""])
    flows.append(_table(rows, [12, 120, 200, 22], font_size=7.2))
    _page_with_flows(c, spec, "確認事項・出典", "99", flows)


# ---------- Excel ----------

def write_xlsx(path, spec, demo, gest, grad=None):
    """見積書として読める Excel。単価・係数・諸経費率を書き換えると金額が再計算される。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    FONT_JP = "游ゴシック"
    ink, sub, line, soft, gold_soft = "333333", "7A7F85", "CDD6DD", "F3F5F6", "F6F1E4"
    thin = Side(style="thin", color=line)
    f_title = Font(name=FONT_JP, size=16, bold=True, color=ink)
    f_h = Font(name=FONT_JP, size=10, bold=True, color="FFFFFF")
    f_b = Font(name=FONT_JP, size=10, color=ink)
    f_bb = Font(name=FONT_JP, size=10, bold=True, color=ink)
    f_s = Font(name=FONT_JP, size=8.5, color=sub)
    fill_h = PatternFill("solid", fgColor=ink)
    fill_cat = PatternFill("solid", fgColor=soft)
    fill_tot = PatternFill("solid", fgColor=gold_soft)
    fill_in = PatternFill("solid", fgColor="FFFDF2")
    YEN = '#,##0;[Red]-#,##0'
    p = spec.get("project", {})
    company = docpages.COMPANY

    def setup(ws, widths, landscape=False):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.sheet_view.showGridLines = False
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.5
        ws.oddFooter.center.text = "&P / &N"
        ws.oddFooter.left.text = "参考資料（推定値を含む）"

    def head(ws, title, ncol):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
        ws["A1"] = title
        ws["A1"].font = f_title
        ws.row_dimensions[1].height = 30
        info = [("物件名", p.get("name", "")), ("所在地", p.get("address", "")), ("作成日", p.get("date", "")),
                ("作成", company or p.get("author", ""))]
        for i, (k, v) in enumerate(info):
            r = 2 + i
            ws.cell(r, 1, k).font = f_s
            ws.cell(r, 2, v).font = f_b
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncol)
        return 7

    def hdr_row(ws, r, labels):
        for i, t in enumerate(labels, 1):
            c = ws.cell(r, i, t)
            c.font, c.fill = f_h, fill_h
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 22

    def box(ws, r1, r2, c1, c2):
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(r, c).border = Border(bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "概要"
    setup(ws, [22, 20, 20, 20, 40])
    r = head(ws, "概算見積（参考）", 5)
    hdr_row(ws, r, ["工事区分", "概算（標準）", "安め", "高め", "主な内容"])
    est_sheets = []
    if demo:
        est_sheets.append(("解体見積", demo, "解体工事"))
    if gest:
        est_sheets.append(("造成見積", gest, "造成工事"))
    first = r + 1
    for i, (sheet, est, label) in enumerate(est_sheets):
        rr = first + i
        ws.cell(rr, 1, label).font = f_bb
        ws.cell(rr, 2, f"='{sheet}'!TOTAL_CELL").number_format = YEN  # 後で置換
        ws.cell(rr, 3, round(est["total"]["low"])).number_format = YEN
        ws.cell(rr, 4, round(est["total"]["high"])).number_format = YEN
        note = ""
        if est is demo:
            q = demo["quantities"]
            note = f"延床 約{q['floor_area_m2']:.0f}㎡（{q['floor_area_tsubo']:.1f}坪）・付帯・残置物・届出"
        else:
            note = f"土工（切土{grad['cut_m3']:.0f}㎥／盛土{grad['fill_m3']:.0f}㎥）・擁壁・排水・調査" if grad else ""
        ws.cell(rr, 5, note).font = f_s
        for c in range(1, 5):
            ws.cell(rr, c).font = f_b if c > 1 else f_bb
    last = first + len(est_sheets) - 1
    tr = last + 1
    ws.cell(tr, 1, "合計（税別）").font = f_bb
    for col in (2, 3, 4):
        L = get_column_letter(col)
        c = ws.cell(tr, col, f"=SUM({L}{first}:{L}{last})")
        c.number_format, c.font = YEN, f_bb
    for c in range(1, 6):
        ws.cell(tr, c).fill = fill_tot
        ws.cell(tr, c).border = Border(top=Side(style="medium", color=ink), bottom=thin)
    box(ws, first, last, 1, 5)
    ws.cell(tr + 1, 1, "消費税（10%）").font = f_b
    ws.cell(tr + 1, 2, f"=ROUND(B{tr}*0.1,0)").number_format = YEN
    ws.cell(tr + 2, 1, "合計（税込）").font = f_bb
    c = ws.cell(tr + 2, 2, f"=B{tr}+B{tr + 1}")
    c.number_format, c.font = YEN, f_bb
    nr = tr + 4
    notes = ["・概算（標準）は各見積シートの合計に連動します（単価・係数・諸経費率を変えると更新）。安め・高めは作成時の参考値です。",
             "・単価は公開相場の標準値です。自社単価がある場合は各見積シートの黄色の単価欄を書き換えてください。"]
    if demo and demo["asbestos_extra"]:
        a = demo["asbestos_extra"][0]
        notes.append(f"・アスベスト除去（含有が判明した場合）は別途：約{a['mid'] / 10000:,.0f}万円（{a['low'] / 10000:,.0f}〜{a['high'] / 10000:,.0f}万円）")
    notes.append("・含まれないもの：地中埋設物の撤去、境界確定測量、建物の設計・確認申請費")
    for n in notes:
        ws.cell(nr, 1, n).font = f_s
        ws.merge_cells(start_row=nr, start_column=1, end_row=nr, end_column=5)
        nr += 1

    # 見積シート
    for sheet, est, label in est_sheets:
        w = wb.create_sheet(sheet)
        setup(w, [44, 9, 6, 12, 7, 14, 40])
        r = head(w, f"{label} 内訳（参考）", 7)
        w.cell(r - 1, 1, "黄色の欄（単価・係数・諸経費率）を書き換えると金額が再計算されます。").font = f_s
        hdr_row(w, r, ["名称", "数量", "単位", "単価", "係数", "金額", "根拠・備考"])
        w.freeze_panes = w.cell(r + 1, 1)
        rf = est.get("region_factor", 1.0)
        cf = est.get("condition_factor", 1.0)
        r += 1
        cats = []
        for ln in est["lines"]:
            if ln["cat"] not in cats:
                cats.append(ln["cat"])
        sub_cells = []
        for cat in cats:
            c = w.cell(r, 1, cat)
            c.font = f_bb
            for col in range(1, 8):
                w.cell(r, col).fill = fill_cat
            r += 1
            start = r
            for ln in [x for x in est["lines"] if x["cat"] == cat]:
                pr = (ln.get("price") or {}).get("mid")
                k = 1.0 if ln["cat"] in ("手続", "調査・設計") else rf * (cf if ln["cat"] in ("建物本体", "付帯工事") else 1.0)
                w.cell(r, 1, ln["name"]).font = f_b
                w.cell(r, 2, ln["qty"]).number_format = "#,##0.0#"
                w.cell(r, 3, ln["unit"])
                c = w.cell(r, 4, pr)
                c.number_format, c.fill = YEN, fill_in
                c = w.cell(r, 5, round(k, 3))
                c.number_format, c.fill = "0.00", fill_in
                w.cell(r, 6, f"=ROUND(B{r}*D{r}*E{r},0)").number_format = YEN
                note = ln.get("basis", "")
                if ln.get("src"):
                    note += f"（出典 {','.join(ln['src'][:3])}）"
                c = w.cell(r, 7, note)
                c.font = f_s
                c.alignment = Alignment(wrap_text=True, vertical="center")
                for col in range(2, 7):
                    w.cell(r, col).font = f_b
                w.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="center")
                for col in range(1, 8):
                    w.cell(r, col).border = Border(bottom=thin)
                lines_needed = max(1, -(-len(ln["name"]) // 22), -(-len(note) // 34))
                w.row_dimensions[r].height = max(18, 14 * lines_needed)
                r += 1
            w.cell(r, 5, "小計").font = f_s
            c = w.cell(r, 6, f"=SUM(F{start}:F{r - 1})")
            c.number_format, c.font = YEN, f_bb
            c.border = Border(top=Side(style="thin", color=sub))
            sub_cells.append(f"F{r}")
            r += 2
        # 合計
        d = r
        w.cell(d, 4, "直接工事費 計").font = f_bb
        w.merge_cells(start_row=d, start_column=4, end_row=d, end_column=5)
        c = w.cell(d, 6, "=" + "+".join(sub_cells))
        c.number_format, c.font = YEN, f_bb
        w.cell(d + 1, 4, "諸経費率").font = f_b
        c = w.cell(d + 1, 5, est["overhead_rate"]["mid"])
        c.number_format, c.fill = "0%", fill_in
        c = w.cell(d + 1, 6, f"=ROUND(F{d}*E{d + 1},0)")
        c.number_format, c.font = YEN, f_b
        w.cell(d + 2, 4, "合計（税別）").font = f_bb
        w.merge_cells(start_row=d + 2, start_column=4, end_row=d + 2, end_column=5)
        c = w.cell(d + 2, 6, f"=F{d}+F{d + 1}")
        c.number_format, c.font = YEN, Font(name=FONT_JP, size=12, bold=True, color=ink)
        for col in range(4, 7):
            w.cell(d + 2, col).fill = fill_tot
            w.cell(d + 2, col).border = Border(top=Side(style="medium", color=ink), bottom=Side(style="medium", color=ink))
        for col in range(1, 8):
            w.cell(d, col).border = Border(top=Side(style="medium", color=ink))
        w.cell(d + 4, 1, f"係数：地域係数 ×{rf}" + (f"、施工条件の割増 ×{cf:.2f}（{'／'.join(est.get('conditions_applied', []))}）" if cf != 1.0 else "")
               + "。手続・調査設計の行は係数1.0。").font = f_s
        # 概要シートの参照を置換
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value == f"='{sheet}'!TOTAL_CELL":
                    cell.value = f"='{sheet}'!F{d + 2}"
        if est is demo and demo["asbestos_extra"]:
            a = demo["asbestos_extra"][0]
            w.cell(d + 5, 1, f"別途（合計に含めない）：{a['name']} 約{round(a['mid']):,}円（{a['basis']}）").font = f_s

    if demo and demo["leftover"]["furniture"]:
        w = wb.create_sheet("残置物")
        setup(w, [30, 10, 10, 16])
        lo = demo["leftover"]
        r = head(w, f"残置物（家財）の推定　{lo['layout']} × {lo['occupied_units']}戸", 4)
        hdr_row(w, r, ["品目", "想定数量", "体積（㎥）", "家電リサイクル"])
        r += 1
        rec = {"fridge": "冷蔵庫", "washer": "洗濯機", "tv": "テレビ", "aircon": "エアコン"}
        s0 = r
        for f in lo["furniture"]:
            w.cell(r, 1, f["name"]).font = f_b
            w.cell(r, 2, f["qty"]).font = f_b
            c = w.cell(r, 3, round(f["m3"], 2))
            c.number_format, c.font = "0.00", f_b
            w.cell(r, 4, rec.get(f.get("recycle"), "")).font = f_s
            for col in range(1, 5):
                w.cell(r, col).border = Border(bottom=thin)
            r += 1
        w.cell(r, 1, "家財の体積 計").font = f_bb
        w.cell(r, 3, f"=SUM(C{s0}:C{r - 1})").number_format = "0.0"
        w.cell(r + 1, 1, f"積込み体積（すき間込み・屋外分を含む）約{lo['volume_m3']}㎥").font = f_s

    # 前提・出典
    w = wb.create_sheet("前提・出典")
    setup(w, [10, 60, 60])
    r = head(w, "前提と出典", 3)
    w.cell(r, 1, "次に確認すること").font = f_bb
    r += 1
    for s_ in spec.get("next_checks", []):
        w.cell(r, 2, "□ " + s_).font = f_b
        r += 1
    r += 1
    hdr_row(w, r, ["ID", "出典", "URL"])
    r += 1
    used = set()
    for est in (demo, gest):
        if est:
            for ln in est["lines"]:
                used.update(ln.get("src") or [])
    seen = set()
    for est in (demo, gest):
        if not est:
            continue
        for s_ in est.get("sources", []):
            if s_["id"] in used and s_["id"] not in seen:
                seen.add(s_["id"])
                w.cell(r, 1, s_["id"]).font = f_s
                w.cell(r, 2, s_.get("title", "")).font = f_b
                c = w.cell(r, 3, s_.get("url", ""))
                c.font = Font(name=FONT_JP, size=9, color="2F5D8A", underline="single")
                if s_.get("url"):
                    c.hyperlink = s_["url"]
                r += 1
    wb.save(path)


# ---------- DXF ----------

class DxfView:
    """elevation.draw_face 等をそのまま DXF に出すための View 互換クラス（単位 mm）。"""

    def __init__(self, msp, ox=0.0, oy=0.0, layer="0"):
        self.msp = msp
        self.ox, self.oy = ox, oy
        self.layer = layer
        self.scale = 1
        self.c = None

    def P(self, x, y):
        return (self.ox + x * 1000, self.oy + y * 1000)

    def line(self, x0, y0, x1, y1, w=0.25, color=None, dash=None):
        self.msp.add_line(self.P(x0, y0), self.P(x1, y1), dxfattribs={"layer": self.layer})

    def polyline(self, pts, w=0.25, color=None, closed=False, fill=None, dash=None):
        if len(pts) >= 2:
            self.msp.add_lwpolyline([self.P(*p) for p in pts], close=closed, dxfattribs={"layer": self.layer})

    def rect(self, x, y, w, h, lw=0.25, color=None, fill=None, dash=None):
        self.polyline([(x, y), (x + w, y), (x + w, y + h), (x, y + h)], closed=True)

    def text(self, x, y, s, size=2.5, anchor="l", color=None, angle=0, dy_mm=0.0):
        t = self.msp.add_text(s, height=size * 40, rotation=angle, dxfattribs={"layer": "TEXT", "style": "JP"})
        t.set_placement(self.P(x, y))

    def hatch(self, *a, **k):
        pass

    def ellipse(self, cx, cy, rx, ry, w=0.25, color=None, dash=None):
        if rx >= ry:
            self.msp.add_ellipse(self.P(cx, cy), major_axis=(rx * 1000, 0), ratio=ry / rx, dxfattribs={"layer": self.layer})
        else:
            self.msp.add_ellipse(self.P(cx, cy), major_axis=(0, ry * 1000), ratio=rx / ry, dxfattribs={"layer": self.layer})

    def len_mm(self, m):
        return m * 1000

    def p(self, x, y):
        return self.P(x, y)

    def dim_h(self, x0, x1, y, off_mm=6.0, label=None, size=2.3):
        d = self.msp.add_linear_dim(base=(self.ox + x0 * 1000, self.oy + y * 1000 + off_mm * 100),
                                    p1=self.P(x0, y), p2=self.P(x1, y), dimstyle="EZDXF",
                                    override={"dimtxt": 150, "dimasz": 80}, dxfattribs={"layer": "DIM"})
        d.render()

    def dim_v(self, y0, y1, x, off_mm=6.0, label=None, size=2.3):
        d = self.msp.add_linear_dim(base=(self.ox + x * 1000 + off_mm * 100, self.oy + y0 * 1000),
                                    p1=self.P(x, y0), p2=self.P(x, y1), angle=90, dimstyle="EZDXF",
                                    override={"dimtxt": 150, "dimasz": 80}, dxfattribs={"layer": "DIM"})
        d.render()

    def level_mark(self, x, y, label, size=2.3, right=True):
        self.text(x, y, "▽" + label, size)


def write_dxf(path, spec, grad):
    import ezdxf

    doc = ezdxf.new("R2010", setup=True)
    doc.styles.add("JP", font="msgothic.ttc")
    for name in ("ELEV", "SITE", "WALL", "CONTOUR", "TEXT", "DIM", "LEVEL"):
        if name not in doc.layers:
            doc.layers.add(name)
    msp = doc.modelspace()
    # 立面図（横に並べる）
    ox = 0.0
    for f in spec.get("building", {}).get("elevations", []):
        v = DxfView(msp, ox, 0, "ELEV")
        elevation.draw_face(v, f)
        elevation.draw_face_dims(v, f)
        v.text(0, -1.5, f["name"], 3)
        x0, x1, _, _ = elevation.face_extent(f)
        ox += (x1 - x0 + 8) * 1000
    # 造成計画平面
    if grad:
        site = spec["site"]
        base_y = -40000
        v = DxfView(msp, 0, base_y, "SITE")
        v.polyline([tuple(p) for p in site["boundary"]], closed=True)
        kbm = grad["road"]["center_avg"] if grad.get("road") else grad["fh_m"]
        vw = DxfView(msp, 0, base_y, "WALL")
        for w in grad["walls"]:
            if w["type"] != "none":
                vw.line(*w["a"], *w["b"])
                vw.text((w["a"][0] + w["b"][0]) / 2, (w["a"][1] + w["b"][1]) / 2, f"{w['type_name']} H={w['h_max']:.2f}", 2)
        vc = DxfView(msp, 0, base_y, "CONTOUR")
        xs = [p[0] for p in site["boundary"]]
        ys = [p[1] for p in site["boundary"]]
        segs = grading_draw._contours(grad["surface"].z, min(xs) - 3, max(xs) + 3, min(ys) - 3, max(ys) + 3, 0.5,
                                      0.25 if grad["existing_diff"] < 2 else 0.5)
        for lv, p, q in segs:
            vc.line(*p, *q)
        vl = DxfView(msp, 0, base_y, "LEVEL")
        step = 3.0
        y = min(ys) + 1
        while y < max(ys):
            x = min(xs) + 1
            while x < max(xs):
                if grading.point_in_poly(x, y, site["boundary"]):
                    vl.text(x, y, f"FH{grad['plan_z'](x, y) - kbm:+.2f} ({grad['surface'].z(x, y) - kbm:+.2f})", 1.6)
                x += step
            y += step
    doc.saveas(path)


# ---------- メイン ----------

def _render_pdf(target, spec, base_dir, demo, grad, gest, total):
    """ページ構成: 表紙 → 概算見積書（総括）→ 内訳書 → 数量の根拠 → 土地の概況 → 写真台帳 → 図面（A3）→ 確認事項・出典"""
    c = rl_canvas.Canvas(target, pagesize=A4)
    c.setTitle(spec.get("project", {}).get("name", "現地写真 調査レポート"))
    c.setAuthor(docpages.COMPANY or "現地写真レポート")
    doc = docpages.Doc(c, spec, total)
    docpages.cover(doc, spec, demo, grad, gest)
    if demo or gest:
        docpages.summary_page(doc, spec, demo, grad, gest)
    if demo:
        docpages.detail_page(doc, demo, "解体工事 内訳書", extra=demo["asbestos_extra"])
    if gest:
        docpages.detail_page(doc, gest, "造成工事 内訳書", notes=docpages.grading_notes(grad))
    if demo:
        docpages.quantities_page(doc, spec, demo)
    docpages.land_page(doc, spec, grad)
    if spec.get("photos"):
        docpages.photo_pages(doc, spec, base_dir, annotate_photo)
    if demo:
        elevation.render_elevations(c, spec)
    if grad:
        grading_draw.render_plan(c, spec, grad, title="現況平面図・解体範囲図（参考）", sheet_no="B-01", mode="existing")
        grading_draw.render_plan(c, spec, grad)
        grading_draw.render_sections(c, spec, grad)
    docpages.notes_page(doc, spec, demo, gest)
    c.save()
    return doc.page

def build(spec, out_dir, base_dir):
    os.makedirs(out_dir, exist_ok=True)
    spec.setdefault("project", {}).setdefault("date", datetime.date.today().isoformat())
    demo = demolition.estimate(spec) if spec.get("building") else None
    grad = gest = None
    if spec.get("site"):
        grad, gest = grading.plan_and_estimate(spec)
    stem = spec.get("project", {}).get("slug", "site_report")
    pdf_path = os.path.join(out_dir, f"{stem}.pdf")
    # 1回目でページ数を数え、2回目に「n / 総数」を入れて描く
    total = _render_pdf(io.BytesIO(), spec, base_dir, demo, grad, gest, None)
    _render_pdf(pdf_path, spec, base_dir, demo, grad, gest, total)
    xlsx_path = os.path.join(out_dir, f"{stem}_見積.xlsx")
    write_xlsx(xlsx_path, spec, demo, gest, grad)
    dxf_path = os.path.join(out_dir, f"{stem}_図面.dxf")
    try:
        write_dxf(dxf_path, spec, grad)
    except Exception as e:  # DXF は補助出力なので失敗してもレポートは出す
        dxf_path = f"(DXF出力失敗: {e})"
    summary = {"pdf": pdf_path, "xlsx": xlsx_path, "dxf": dxf_path}
    if demo:
        summary["demolition_total"] = {k: round(v) for k, v in demo["total"].items()}
        summary["floor_area_m2"] = round(demo["quantities"]["floor_area_m2"], 1)
        summary["floor_area_tsubo"] = round(demo["quantities"]["floor_area_tsubo"], 1)
        summary["per_tsubo_mid"] = round(demo["per_tsubo_mid"] or 0)
        summary["asbestos_extra"] = [{k: (round(v) if isinstance(v, float) else v) for k, v in a.items() if k in ("name", "low", "mid", "high")}
                                     for a in demo["asbestos_extra"]]
        summary["demolition_lines"] = [{"cat": l["cat"], "name": l["name"], "qty": l["qty"], "unit": l["unit"],
                                        "mid": round(l["mid"]) if l["mid"] is not None else None} for l in demo["lines"]]
    if grad:
        summary["grading_total"] = {k: round(v) for k, v in gest["total"].items()}
        summary["grading"] = {k: grad[k] for k in ("area_m2", "fh_m", "fh_basis", "cut_m3", "fill_m3", "existing_diff", "flags")}
        kbm = grad["road"]["center_avg"] if grad.get("road") else 0.0
        summary["kbm_tp"] = round(kbm, 2)
        summary["alternatives"] = [{**a, "fh_rel": round(a["fh_m"] - kbm, 2), "total_mid": round(a["total_mid"])}
                                   for a in grad.get("alternatives", [])]
        summary["grading_lines"] = [{"cat": l["cat"], "name": l["name"], "qty": l["qty"], "unit": l["unit"],
                                     "mid": round(l["mid"]) if l["mid"] is not None else None} for l in gest["lines"]]
        summary["walls"] = [{k: w[k] for k in ("index", "length_m", "h_max", "type_name", "side")} for w in grad["walls"]]
    return summary
